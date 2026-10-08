# -*- coding: utf-8 -*-
"""
Actualiza únicamente los Excel de predicciones 30/60/90 días en Google Drive.

Este archivo es independiente de 02_actualizar_observado_hidromet.py:
- no modifica ese script;
- no ejecuta DWLT;
- conserva el libro Excel descargado, su hoja, columnas, estilos y marcas;
- consulta HidroMet solo para la ventana reciente;
- actualiza el mismo archivo de Drive por su file_id.

Variables de entorno:
    PRED_DRIVE_FOLDER_ID   ID de la carpeta PREDICCIONES306090.
    GOOGLE_SERVICE_JSON    JSON de la cuenta de servicio (también acepta
                           PRED_GOOGLE_SERVICE_JSON).
    PRED_API_START_DATE   YYYY-MM-DD; por defecto, hoy - 7 días.
    PRED_API_END_DATE     YYYY-MM-DD; por defecto, hoy.
    PRED_DRIVE_WRITE      true/false; por defecto true. Use false para ensayo.

La cuenta de servicio debe tener permiso Editor sobre la carpeta de Drive.
"""

from __future__ import annotations

import copy
import importlib.util
import os
import tempfile
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent.parent
BASE_SCRIPT = BASE_DIR / "scripts" / "02_actualizar_observado_hidromet.py"

# Carga las funciones de consulta y Drive sin ejecutar su main().
spec = importlib.util.spec_from_file_location("hidromet_base", BASE_SCRIPT)
if spec is None or spec.loader is None:
    raise RuntimeError(f"No se pudo cargar {BASE_SCRIPT}")
hidromet = importlib.util.module_from_spec(spec)
spec.loader.exec_module(hidromet)

PRED_DRIVE_FOLDER_ID = os.getenv(
    "PRED_DRIVE_FOLDER_ID",
    os.getenv("OBS_DRIVE_FOLDER_ID", ""),
).strip()
hidromet.GOOGLE_SERVICE_JSON = os.getenv(
    "PRED_GOOGLE_SERVICE_JSON",
    os.getenv("GOOGLE_SERVICE_JSON", ""),
).strip()

def _fecha_defecto_inicio() -> str:
    return (date.today() - timedelta(days=7)).strftime("%Y-%m-%d")


PRED_API_START_DATE = os.getenv(
    "PRED_API_START_DATE",
    _fecha_defecto_inicio(),
).strip()
PRED_API_END_DATE = os.getenv(
    "PRED_API_END_DATE",
    date.today().strftime("%Y-%m-%d"),
).strip()
PRED_DRIVE_WRITE = os.getenv("PRED_DRIVE_WRITE", "true").strip().lower() in {
    "1", "true", "yes", "si", "sí",
}

PRED_BACKEND_DIR = BASE_DIR / "backend" / "predicciones_30_60_90"
PRED_CACHE_DIR = PRED_BACKEND_DIR / "cache"
PRED_INPUT_PARQUET = PRED_CACHE_DIR / "niveles_drive_predicciones.parquet"
PRED_INPUT_CSV = PRED_CACHE_DIR / "niveles_drive_predicciones.csv"
PRED_MANIFEST_CSV = PRED_CACHE_DIR / "manifiesto_actualizacion_drive.csv"

ESTACIONES_PREDICCION = {
    "BELLAVISTA",
    "BORJA",
    "CONTAMANA",
    "ENAPU",
    "LAGUNAS",
    "NAUTA",
    "PUERTO ALEGRIA",
    "REQUENA",
    "SAN REGIS",
    "SANTA MARIA DE NANAY",
    "TAMSHIYACU",
    "YURIMAGUAS",
}

COLUMNAS_REQUERIDAS = [
    "ID_ESTACION",
    "ESTACION_NOMBRE",
    "FECHA",
    "ANIO",
    "MES",
    "DIA",
    "H6",
    "H10",
    "H14",
    "H18",
    "H_PROM",
    "FECHA_API",
    "FECHA_CONSULTA_INI",
    "FECHA_CONSULTA_FIN",
    "ESTADO_FECHA",
    "Id",
    "Estacion",
]


def clave_estacion(nombre: str) -> str:
    return hidromet.normalizar_texto(Path(str(nombre)).stem)


def _valor_fecha(valor: Any) -> date | None:
    if valor is None or (isinstance(valor, float) and np.isnan(valor)):
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    ts = pd.to_datetime(valor, errors="coerce")
    if pd.isna(ts):
        return None
    return ts.date()


def _mapa_encabezados(ws) -> dict[str, int]:
    mapa = {}
    for cell in ws[1]:
        if cell.value is not None:
            mapa[str(cell.value).strip().upper()] = cell.column
    return mapa


def _copiar_estilo_fila(ws, fila_origen: int, fila_destino: int) -> None:
    if fila_origen < 2 or fila_destino < 2:
        return
    for col in range(1, ws.max_column + 1):
        origen = ws.cell(fila_origen, col)
        destino = ws.cell(fila_destino, col)
        if origen.has_style:
            destino._style = copy.copy(origen._style)
        if origen.number_format:
            destino.number_format = origen.number_format
        if origen.alignment:
            destino.alignment = copy.copy(origen.alignment)
        if origen.protection:
            destino.protection = copy.copy(origen.protection)
    if fila_origen in ws.row_dimensions:
        ws.row_dimensions[fila_destino].height = ws.row_dimensions[fila_origen].height


def _id_desde_libro(ws, headers: dict[str, int], nombre: str) -> int | None:
    for etiqueta in ("ID_ESTACION", "ID", "ID ESTACION"):
        col = headers.get(etiqueta)
        if col is None:
            continue
        for fila in range(2, min(ws.max_row, 200) + 1):
            valor = ws.cell(fila, col).value
            numero = pd.to_numeric(valor, errors="coerce")
            if pd.notna(numero):
                return int(numero)
    mapa = {
        nombre: int(id_estacion)
        for id_estacion, nombre in hidromet.ESTACIONES_API.items()
    }
    return mapa.get(nombre)


def _fila_por_fecha(ws, headers: dict[str, int]) -> dict[date, int]:
    col_fecha = headers.get("FECHA")
    if col_fecha is None:
        return {}
    resultado = {}
    for fila in range(2, ws.max_row + 1):
        fecha = _valor_fecha(ws.cell(fila, col_fecha).value)
        if fecha is not None:
            resultado.setdefault(fecha, fila)
    return resultado


def _poner(ws, headers: dict[str, int], fila: int, columna: str, valor: Any) -> None:
    col = headers.get(columna.upper())
    if col is not None:
        ws.cell(fila, col).value = valor


def actualizar_libro_prediccion(
    path_excel: Path,
    df_api: pd.DataFrame,
    id_estacion: int,
    estacion_nombre: str,
    fecha_ini: str,
    fecha_fin: str,
) -> dict[str, int]:
    """Actualiza datos API en una copia local del libro, conservando estructura y estilos."""
    wb = load_workbook(path_excel, keep_links=True)
    ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb[wb.sheetnames[0]]
    headers = _mapa_encabezados(ws)

    faltantes = [c for c in COLUMNAS_REQUERIDAS if c.upper() not in headers]
    if faltantes:
        raise ValueError(f"{path_excel.name}: faltan columnas {faltantes}")

    df = df_api.copy()
    if df.empty:
        return {"actualizadas": 0, "insertadas": 0}

    df["FECHA"] = pd.to_datetime(df["FECHA"], errors="coerce").dt.normalize()
    df["H_PROM"] = pd.to_numeric(df["H_PROM"], errors="coerce")
    df = df.dropna(subset=["FECHA", "H_PROM"]).copy()
    df = df[(df["H_PROM"] > 0) & (df["H_PROM"] != -999)].copy()
    df = df.sort_values("FECHA").drop_duplicates("FECHA", keep="last")

    filas = _fila_por_fecha(ws, headers)
    actualizadas = 0
    insertadas = 0

    for _, registro in df.iterrows():
        fecha_ts = registro["FECHA"]
        fecha = fecha_ts.date()
        fila = filas.get(fecha)

        if fila is None:
            fila = ws.max_row + 1
            _copiar_estilo_fila(ws, fila - 1, fila)
            filas[fecha] = fila
            insertadas += 1
        else:
            actualizadas += 1

        valores = {
            "ID_ESTACION": id_estacion,
            "ESTACION_NOMBRE": estacion_nombre,
            "FECHA": datetime(fecha.year, fecha.month, fecha.day),
            "ANIO": fecha.year,
            "MES": fecha.month,
            "DIA": fecha.day,
            "H6": registro.get("H6"),
            "H10": registro.get("H10"),
            "H14": registro.get("H14"),
            "H18": registro.get("H18"),
            "H_PROM": registro.get("H_PROM"),
            "FECHA_API": registro.get("FECHA_API", ""),
            "FECHA_CONSULTA_INI": fecha_ini,
            "FECHA_CONSULTA_FIN": fecha_fin,
            "ESTADO_FECHA": "HIDROMET_API",
            "Id": id_estacion,
            "Estacion": estacion_nombre,
        }

        for columna, valor in valores.items():
            if pd.isna(valor):
                continue
            _poner(ws, headers, fila, columna, valor)

    wb.save(path_excel)
    return {"actualizadas": actualizadas, "insertadas": insertadas}


def _nombre_drive_por_clave(indice: dict[str, dict], nombre: str) -> dict | None:
    return indice.get(clave_estacion(nombre))


def _leer_libro_backend(path_excel: Path, estacion: str) -> pd.DataFrame:
    """Lee un Excel descargado para alimentar el backend propio de predicciones."""
    try:
        df = pd.read_excel(path_excel, sheet_name="Sheet1")
    except Exception:
        df = pd.read_excel(path_excel)

    df.columns = [str(c).strip() for c in df.columns]
    if "FECHA" not in df.columns or "H_PROM" not in df.columns:
        return pd.DataFrame()

    salida = pd.DataFrame()
    salida["estacion"] = estacion
    salida["id_estacion"] = pd.to_numeric(df.get("ID_ESTACION"), errors="coerce")
    salida["fecha"] = pd.to_datetime(df["FECHA"], errors="coerce").dt.normalize()
    columnas_horarias = []
    for origen, destino in (("H6", "h6"), ("H10", "h10"), ("H14", "h14"), ("H18", "h18")):
        salida[destino] = pd.to_numeric(df.get(origen), errors="coerce")
        columnas_horarias.append(destino)
    nivel_hprom = pd.to_numeric(df["H_PROM"], errors="coerce")
    nivel_horario = salida[columnas_horarias].replace(-999, np.nan).mean(axis=1, skipna=True)
    salida["nivel_m"] = nivel_hprom.fillna(nivel_horario)
    estado = df.get("ESTADO_FECHA")
    salida["estado_fecha"] = estado.astype(str) if estado is not None else ""
    salida["fuente"] = "PREDICCIONES306090_DRIVE"
    salida = salida.dropna(subset=["fecha"]).sort_values(["estacion", "fecha"])
    print(
        f"[CACHE] {estacion}: filas={len(salida):,} | "
        f"niveles_validos={salida['nivel_m'].notna().sum():,}"
    )
    return salida.reset_index(drop=True)


def _guardar_cache_backend(libros: dict[str, pd.DataFrame]) -> None:
    partes = [df for df in libros.values() if df is not None and not df.empty]
    if not partes:
        print("[CACHE] No se generó caché: no hubo libros legibles.")
        return
    PRED_CACHE_DIR.mkdir(parents=True, exist_ok=True)
    cache = pd.concat(partes, ignore_index=True)
    cache = cache.drop_duplicates(["estacion", "fecha"], keep="last")
    cache = cache.sort_values(["estacion", "fecha"]).reset_index(drop=True)
    print(
        f"[CACHE] fechas_validas={cache['fecha'].notna().sum():,} | "
        f"rango={cache['fecha'].min()} a {cache['fecha'].max()} | "
        f"tipo_fecha={cache['fecha'].dtype}"
    )
    cache.to_parquet(PRED_INPUT_PARQUET, index=False)
    cache.to_csv(PRED_INPUT_CSV, index=False, encoding="utf-8-sig")
    print(f"[CACHE] {PRED_INPUT_PARQUET} | filas: {len(cache):,}")
    print(f"[CACHE] {PRED_INPUT_CSV}")


def procesar_drive() -> int:
    if not PRED_DRIVE_FOLDER_ID:
        raise RuntimeError("Falta PRED_DRIVE_FOLDER_ID (ID de la carpeta PREDICCIONES306090).")

    service = hidromet.construir_drive_service()
    if service is None:
        raise RuntimeError("No se pudo crear el servicio Drive. Revise GOOGLE_SERVICE_JSON.")

    archivos = hidromet.listar_archivos_drive(service, PRED_DRIVE_FOLDER_ID)
    indice = hidromet.construir_indice_drive_por_estacion(archivos)

    print(f"Carpeta Drive: {PRED_DRIVE_FOLDER_ID}")
    print(f"Archivos Excel detectados: {len(indice)}")
    print(f"Ventana HidroMet: {PRED_API_START_DATE} a {PRED_API_END_DATE}")
    print(f"Modo escritura: {'ACTIVO' if PRED_DRIVE_WRITE else 'ENSAYO (sin subir cambios)'}")

    fallback_ids = {
        hidromet.normalizar_texto(nombre): int(id_estacion)
        for id_estacion, nombre in hidromet.ESTACIONES_API.items()
    }

    resumen = []
    libros_backend = {}
    with tempfile.TemporaryDirectory(prefix="amaru_pred_drive_") as tmp:
        carpeta_tmp = Path(tmp)

        for nombre in sorted(ESTACIONES_PREDICCION):
            archivo = _nombre_drive_por_clave(indice, nombre)
            if archivo is None:
                resumen.append((nombre, "OMITIDO_SIN_ARCHIVO", 0, 0))
                print(f"[OMITIDO] {nombre}: no existe un Excel con ese nombre.")
                continue

            file_id = archivo.get("id")
            nombre_drive = archivo.get("name", "")
            if not file_id or not nombre_drive.lower().endswith(".xlsx"):
                resumen.append((nombre, "OMITIDO_FORMATO", 0, 0))
                print(f"[OMITIDO] {nombre_drive}: solo se aceptan .xlsx.")
                continue

            path = hidromet.descargar_archivo_drive(
                service, file_id, nombre_drive, carpeta_tmp
            )
            if path is None:
                resumen.append((nombre, "ERROR_DESCARGA", 0, 0))
                continue

            wb = load_workbook(path, read_only=False, keep_links=True)
            ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb[wb.sheetnames[0]]
            headers = _mapa_encabezados(ws)
            id_estacion = _id_desde_libro(ws, headers, nombre)
            wb.close()
            libro_inicial = _leer_libro_backend(path, nombre)
            if not libro_inicial.empty:
                libros_backend[nombre] = libro_inicial

            if id_estacion is None:
                id_estacion = fallback_ids.get(nombre)

            if id_estacion is None:
                resumen.append((nombre, "OMITIDO_SIN_ID_API", 0, 0))
                print(f"[OMITIDO] {nombre}: no se encontró ID_ESTACION.")
                continue

            print(f"[API] {nombre} | ID {id_estacion}")
            df_api = hidromet.descargar_api_hidromet_formato_excel_estacion(
                id_estacion=id_estacion,
                estacion_nombre=nombre,
                fecha_ini=PRED_API_START_DATE,
                fecha_fin=PRED_API_END_DATE,
            )

            if df_api.empty:
                resumen.append((nombre, "SIN_DATOS_API", 0, 0))
                print(f"[SIN DATOS] {nombre}")
                continue

            conteo = actualizar_libro_prediccion(
                path_excel=path,
                df_api=df_api,
                id_estacion=id_estacion,
                estacion_nombre=nombre,
                fecha_ini=PRED_API_START_DATE,
                fecha_fin=PRED_API_END_DATE,
            )

            libro_actualizado = _leer_libro_backend(path, nombre)
            if not libro_actualizado.empty:
                libros_backend[nombre] = libro_actualizado

            if PRED_DRIVE_WRITE:
                ok = hidromet.actualizar_archivo_excel_drive(service, file_id, path)
                estado = "ACTUALIZADO" if ok else "ERROR_SUBIDA"
            else:
                estado = "ENSAYO"

            resumen.append((nombre, estado, conteo["actualizadas"], conteo["insertadas"]))
            print(
                f"[{estado}] {nombre}: "
                f"{conteo['actualizadas']} actualizadas, "
                f"{conteo['insertadas']} insertadas"
            )

    _guardar_cache_backend(libros_backend)
    print("\nResumen PREDICCIONES306090")
    for nombre, estado, actualizadas, insertadas in resumen:
        print(f"- {nombre}: {estado} | existentes={actualizadas} | nuevas={insertadas}")

    return 0 if not any(e.startswith("ERROR") for _, e, _, _ in resumen) else 1


if __name__ == "__main__":
    raise SystemExit(procesar_drive())


